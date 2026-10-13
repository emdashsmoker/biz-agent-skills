#!/usr/bin/env python3
"""Check a project plan (plan.json) and export it to Markdown, CSV and XLSX.

Why this script exists: dates, buffer ratios and M/M sums must be computed by code,
not by a language model, so that anyone can reproduce the same numbers.

Usage:
    python plan.py check plan.json
    python plan.py md    plan.json -o overview.md
    python plan.py csv   plan.json -o wbs.csv
    python plan.py xlsx  plan.json -o wbs.xlsx

Only the standard library is needed for check, md and csv. xlsx needs openpyxl.
"""
import argparse
import csv
import json
import sys
from datetime import date, timedelta

# ---------------------------------------------------------------------------
# Settings. Change DEFAULT_LANG to "en" for English labels.
# ---------------------------------------------------------------------------
DEFAULT_LANG = "ko"
DEFAULT_BUFFER_RATIO = 0.15      # slack before contract end, share of total period
DEFAULT_MAX_TASK_DAYS = 28       # a work row longer than this should be split
ADVISORY_GAP_DAYS = 28           # minimum gap between advisory meetings
ADVISORY_BEFORE_FINAL_DAYS = 28  # last advisory meeting before the final report
HEARING_BEFORE_FINAL_DAYS = 42   # public hearing before the final report
HEARING_CRITICAL_DAYS = 28       # closer than this leaves no time to reflect
MAX_SUCCESS_LINES = 3
MM_TOLERANCE = 0.01
MONTH_MM_LIMIT = 1.0             # one person cannot exceed 1 M/M in a month
WEEKLY_GANTT_MAX_WEEKS = 30      # about 7 months; longer projects get monthly columns
DAYS_PER_MONTH = 365.25 / 12

TYPES = ("work", "wait", "check", "buffer")
STATUSES = ("planned", "doing", "done")
EVENTS = ("kickoff", "interim_report", "final_report", "advisory", "hearing", "other")
RACI_LETTERS = ("R", "A", "C", "I")
SYMBOL = {"work": "■", "wait": "○", "check": "●", "buffer": "□"}

LABELS = {
    "ko": {
        "type": {"work": "작업", "wait": "대기", "check": "점검", "buffer": "여유"},
        "status": {"planned": "예정", "doing": "진행", "done": "완료"},
        "wbs_cols": ["No", "단계", "작업", "구분", "산출물(완료 기준)", "담당",
                     "시작", "종료", "기간(일)", "상태", "비고"],
        "mm_cols": ["No", "성명", "역할", "참여율", "계약 M/M"],
        "mm_sum": "합계", "mm_diff": "차이", "mm_month_sum": "월 합계",
        "sheet_wbs": "WBS", "sheet_mm": "투입인력",
        "period": "기간", "legend": "■ 작업   ○ 대기   ● 점검   □ 여유",
        "week_note": "   (간트 머리글의 숫자는 각 주 월요일의 날짜)",
        "week_label": "{y}년 {m}월", "year_label": "{y}년", "month_fmt": 'm"월"',
        "font": "맑은 고딕",
        "h_basic": "기본개요", "l_name": "프로젝트명", "l_period": "프로젝트 기간",
        "h_why": "Why(수혜자 & 니즈 및 문제)", "l_benef": "수혜자", "l_needs": "니즈 및 문제",
        "h_goals": "프로젝트 목표", "l_goal": "목표{n}",
        "h_wbs": "WBS", "h_hyp": "가설 & 솔루션", "l_hyp": "가설 {n} : {text}",
        "h_success": "성공판단기준", "h_checks": "보고/자문회의 등 프로젝트 점검 일정",
        "h_notes": "특기사항", "h_detail": "WBS 상세", "h_mm": "투입인력",
        "l_staff": "투입 인력", "staff_line": "{n}명, 총 {mm} M/M",
        "months": "{n}개월",
        "buffer_line": "- (일정 여유) 마지막 작업 종료 {last}, 계약 종료 {end}까지 {days}일 확보 (전체 기간의 {pct}%)",
        "month": "{m}월", "year_month": "{y}년 {m}월",
        "msg": {
            "err_field": "필수 항목 누락: {where} {field}",
            "err_date": "날짜 형식 오류(YYYY-MM-DD): {where} {value}",
            "err_order": "종료일이 시작일보다 빠름: {where}",
            "err_type": "구분 값 오류: {where} {value} (work, wait, check, buffer 중 선택)",
            "err_status": "상태 값 오류: {where} {value} (planned, doing, done 중 선택)",
            "err_event": "event 값 오류: {where} {value}",
            "w_outside": "계약 기간 밖의 일정: {where}",
            "w_owner": "담당 미지정: {where}",
            "w_output": "산출물(완료 기준) 미기재: {where}",
            "w_long": "작업 기간 {days}일로 {limit}일 초과, 중간 산출물 기준 분할 검토: {where}",
            "w_buffer": "일정 여유 {days}일(전체의 {pct}%)로 기준 {target}% 미달",
            "w_buffer_row": "일정 여유 구간이 표에 없음, 구분 buffer 행 추가 권장",
            "w_adv_gap": "자문회의 간격 {days}일로 {limit}일 미만: {a} / {b}",
            "w_adv_final": "마지막 자문회의와 최종보고 간격 {days}일로 {limit}일 미만",
            "w_hear_final": "공청회와 최종보고 간격 {days}일로 {limit}일 미만",
            "w_hear_crit": "공청회와 최종보고 간격 {days}일, 의견 반영 기간 부족으로 납기 직결",
            "w_success": "성공판단기준 {n}줄로 {limit}줄 초과",
            "w_mm_total": "M/M 합계 불일치: {name} 배분 {got} / 계약 {want}",
            "w_mm_month": "월 M/M {limit} 초과: {name} {month} {value}",
            "w_mm_outside": "계약 기간 밖의 M/M: {name} {month}",
            "w_raci_a": "RACI의 A가 {n}명(1명 필요): {where}",
            "w_raci_r": "RACI의 R 미지정: {where}",
            "w_raci_none": "RACI 미기재 작업 {n}행: {rows} 등",
            "w_progress": "일정 여유 소진 {used}%가 진척 {done}%보다 큼, 대응 경로 검토 필요",
            "sum_head": "[요약]",
            "sum_period": "계약 기간 {start} - {end} ({days}일)",
            "sum_rows": "작업 행 {n}개, 마지막 작업 종료 {last}",
            "sum_buffer": "일정 여유 {days}일 (전체의 {pct}%, 기준 {target}%)",
            "sum_progress": "진척 {done}%, 일정 여유 소진 {used}%",
            "sum_mm": "M/M {name}: 배분 {got} / 계약 {want}",
            "sum_ok": "경고 없음",
            "sum_counts": "오류 {e}건, 경고 {w}건",
        },
    },
    "en": {
        "type": {"work": "Work", "wait": "Wait", "check": "Checkpoint", "buffer": "Buffer"},
        "status": {"planned": "Planned", "doing": "In progress", "done": "Done"},
        "wbs_cols": ["No", "Stage", "Task", "Type", "Deliverable (done when)", "Owner",
                     "Start", "End", "Days", "Status", "Notes"],
        "mm_cols": ["No", "Name", "Role", "Allocation", "Contract M/M"],
        "mm_sum": "Total", "mm_diff": "Diff", "mm_month_sum": "Monthly total",
        "sheet_wbs": "WBS", "sheet_mm": "Staffing",
        "period": "Period", "legend": "■ Work   ○ Wait   ● Checkpoint   □ Buffer",
        "week_note": "   (Gantt header numbers are the Monday of each week)",
        "week_label": "{mon} {y}", "year_label": "{y}", "month_fmt": "mmm",
        "font": "Calibri",
        "h_basic": "Basic Information", "l_name": "Project name", "l_period": "Project period",
        "h_why": "Why (Beneficiaries & Needs/Problems)", "l_benef": "Beneficiaries",
        "l_needs": "Needs and problems",
        "h_goals": "Project Goals", "l_goal": "Goal {n}",
        "h_wbs": "WBS", "h_hyp": "Hypotheses & Solutions", "l_hyp": "Hypothesis {n} : {text}",
        "h_success": "Success Criteria", "h_checks": "Reviews, Advisory Meetings and Other Checkpoints",
        "h_notes": "Notes", "h_detail": "WBS Detail", "h_mm": "Staffing",
        "l_staff": "Staffing", "staff_line": "{n} people, {mm} M/M in total",
        "months": "{n} months",
        "buffer_line": "- (Schedule buffer) Last task ends {last}, leaving {days} days until the contract end on {end} ({pct}% of the total period)",
        "month": "{mon}", "year_month": "{mon} {y}",
        "msg": {
            "err_field": "Missing field: {where} {field}",
            "err_date": "Invalid date (YYYY-MM-DD): {where} {value}",
            "err_order": "End date before start date: {where}",
            "err_type": "Invalid type: {where} {value} (use work, wait, check, buffer)",
            "err_status": "Invalid status: {where} {value} (use planned, doing, done)",
            "err_event": "Invalid event: {where} {value}",
            "w_outside": "Dates outside the contract period: {where}",
            "w_owner": "No owner: {where}",
            "w_output": "No deliverable (done when): {where}",
            "w_long": "Task lasts {days} days, over {limit}; consider splitting by interim output: {where}",
            "w_buffer": "Schedule buffer {days} days ({pct}% of total) is below the {target}% target",
            "w_buffer_row": "No buffer row in the table; add a row with type buffer",
            "w_adv_gap": "Advisory meetings only {days} days apart (min {limit}): {a} / {b}",
            "w_adv_final": "Last advisory meeting is {days} days before the final report (min {limit})",
            "w_hear_final": "Public hearing is {days} days before the final report (min {limit})",
            "w_hear_crit": "Public hearing is {days} days before the final report; too little time to reflect input, deadline at risk",
            "w_success": "Success criteria have {n} lines (max {limit})",
            "w_mm_total": "M/M total mismatch: {name} allocated {got} / contract {want}",
            "w_mm_month": "Monthly M/M over {limit}: {name} {month} {value}",
            "w_mm_outside": "M/M outside the contract period: {name} {month}",
            "w_raci_a": "RACI has {n} A (need exactly 1): {where}",
            "w_raci_r": "RACI has no R: {where}",
            "w_raci_none": "{n} work rows have no RACI: {rows} and others",
            "w_progress": "Buffer used {used}% exceeds progress {done}%; review recovery options",
            "sum_head": "[Summary]",
            "sum_period": "Contract period {start} - {end} ({days} days)",
            "sum_rows": "{n} rows, last task ends {last}",
            "sum_buffer": "Schedule buffer {days} days ({pct}% of total, target {target}%)",
            "sum_progress": "Progress {done}%, buffer used {used}%",
            "sum_mm": "M/M {name}: allocated {got} / contract {want}",
            "sum_ok": "No warnings",
            "sum_counts": "{e} errors, {w} warnings",
        },
    },
}
MONTH_ABBR = ("Jan", "Feb", "Mar", "Apr", "May", "Jun",
              "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------
def parse_date(value):
    """Return a date for 'YYYY-MM-DD' or None when the value is not valid."""
    try:
        return date.fromisoformat(str(value))
    except ValueError:
        return None


def days_between(start, end):
    """Inclusive day count, so a one-day task has duration 1."""
    return (end - start).days + 1


def pct(part, whole):
    return round(100 * part / whole, 1) if whole else 0.0


def month_keys(start, end):
    """List 'YYYY-MM' keys for every calendar month touched by the period."""
    keys, y, m = [], start.year, start.month
    while (y, m) <= (end.year, end.month):
        keys.append(f"{y:04d}-{m:02d}")
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return keys


def month_label(lang, d, with_year):
    lab = LABELS[lang]
    key = "year_month" if with_year else "month"
    return lab[key].format(y=d.year, m=d.month, mon=MONTH_ABBR[d.month - 1])


def span_label(lang, start, end):
    """Month span such as '3월-4월' for stage headers."""
    cross_year = start.year != end.year
    first = month_label(lang, start, cross_year)
    last = month_label(lang, end, cross_year)
    return first if (start.year, start.month) == (end.year, end.month) else f"{first}-{last}"


def months_text(lang, start, end):
    n = days_between(start, end) / DAYS_PER_MONTH
    shown = str(round(n)) if abs(n - round(n)) < 0.1 else f"{n:.1f}"
    return LABELS[lang]["months"].format(n=shown)


def task_where(task, index):
    return f"#{index + 1} {task.get('name', '')}"


# ---------------------------------------------------------------------------
# Loading and validation
# ---------------------------------------------------------------------------
class Report:
    """Collects errors (data cannot be used) and warnings (plan needs a look)."""

    def __init__(self, lang):
        self.msg = LABELS[lang]["msg"]
        self.errors, self.warnings, self.summary = [], [], []

    def error(self, key, **kw):
        self.errors.append(self.msg[key].format(**kw))

    def warn(self, key, **kw):
        self.warnings.append(self.msg[key].format(**kw))

    def note(self, key, **kw):
        self.summary.append(self.msg[key].format(**kw))


def load_plan(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def check_project(plan, rep):
    proj = plan.get("project", {})
    for field in ("name", "start", "end"):
        if not proj.get(field):
            rep.error("err_field", where="project", field=field)
    start, end = parse_date(proj.get("start")), parse_date(proj.get("end"))
    if proj.get("start") and not start:
        rep.error("err_date", where="project.start", value=proj.get("start"))
    if proj.get("end") and not end:
        rep.error("err_date", where="project.end", value=proj.get("end"))
    if start and end and end < start:
        rep.error("err_order", where="project")
    return start, end


def check_tasks(plan, rep, p_start, p_end):
    """Validate each row and return rows with parsed dates for later checks."""
    rows = []
    raci_on = plan.get("options", {}).get("raci", False)
    max_days = plan.get("project", {}).get("max_task_days", DEFAULT_MAX_TASK_DAYS)
    missing_raci = []
    for i, task in enumerate(plan.get("tasks", [])):
        where = task_where(task, i)
        for field in ("stage", "name", "type", "start", "end"):
            if not task.get(field):
                rep.error("err_field", where=where, field=field)
        kind = task.get("type")
        if kind and kind not in TYPES:
            rep.error("err_type", where=where, value=kind)
        status = task.get("status", "planned")
        if status not in STATUSES:
            rep.error("err_status", where=where, value=status)
        event = task.get("event")
        if event and event not in EVENTS:
            rep.error("err_event", where=where, value=event)
        start, end = parse_date(task.get("start")), parse_date(task.get("end"))
        if task.get("start") and not start:
            rep.error("err_date", where=where, value=task.get("start"))
        if task.get("end") and not end:
            rep.error("err_date", where=where, value=task.get("end"))
        if not (start and end):
            continue
        if end < start:
            rep.error("err_order", where=where)
            continue
        if p_start and p_end and (start < p_start or end > p_end):
            rep.warn("w_outside", where=where)
        if kind in ("work", "check") and not task.get("owner"):
            rep.warn("w_owner", where=where)
        if kind in ("work", "check") and not task.get("output"):
            rep.warn("w_output", where=where)
        if kind == "work" and not task.get("allow_long") and days_between(start, end) > max_days:
            rep.warn("w_long", days=days_between(start, end), limit=max_days, where=where)
        if raci_on and kind == "work":
            if task.get("raci"):
                check_raci_row(task["raci"], where, rep)
            else:
                missing_raci.append(where)
        rows.append(dict(task, _start=start, _end=end, _status=status))
    if missing_raci:
        rep.warn("w_raci_none", n=len(missing_raci), rows=", ".join(missing_raci[:3]))
    return rows


def check_raci_row(raci, where, rep):
    letters = list(raci.values())
    if letters.count("A") != 1:
        rep.warn("w_raci_a", n=letters.count("A"), where=where)
    if "R" not in letters:
        rep.warn("w_raci_r", where=where)


def check_buffer(plan, rows, rep, p_start, p_end):
    """Slack between the last work or wait row and the contract end date."""
    real = [r for r in rows if r.get("type") in ("work", "wait")]
    if not (real and p_start and p_end):
        return None
    total = days_between(p_start, p_end)
    last = max(r["_end"] for r in real)
    slack = max((p_end - last).days, 0)
    target = plan.get("project", {}).get("buffer_ratio", DEFAULT_BUFFER_RATIO)
    rep.note("sum_period", start=p_start, end=p_end, days=total)
    rep.note("sum_rows", n=len(rows), last=last)
    rep.note("sum_buffer", days=slack, pct=pct(slack, total), target=round(target * 100))
    if slack < total * target:
        rep.warn("w_buffer", days=slack, pct=pct(slack, total), target=round(target * 100))
    if slack > 0 and not any(r.get("type") == "buffer" for r in rows):
        rep.warn("w_buffer_row")
    return {"last": last, "slack": slack, "total": total}


def check_progress(plan, rows, rep, buffer_info):
    """Compare buffer used with work done, a simplified buffer fever chart."""
    baseline = plan.get("project", {}).get("baseline_last_end")
    base_last = parse_date(baseline) if baseline else None
    if not (base_last and buffer_info):
        return
    p_end = parse_date(plan["project"]["end"])
    base_slack = max((p_end - base_last).days, 0)
    used = max(base_slack - buffer_info["slack"], 0)
    used_pct = pct(used, base_slack) if base_slack else 100.0
    work = [r for r in rows if r.get("type") == "work"]
    total_days = sum(days_between(r["_start"], r["_end"]) for r in work)
    done_days = sum(days_between(r["_start"], r["_end"]) for r in work if r["_status"] == "done")
    done_pct = pct(done_days, total_days)
    rep.note("sum_progress", done=done_pct, used=used_pct)
    if used_pct > done_pct:
        rep.warn("w_progress", used=used_pct, done=done_pct)


def check_events(rows, rep):
    by_event = {}
    for r in rows:
        if r.get("event"):
            by_event.setdefault(r["event"], []).append(r)
    advisory = sorted(by_event.get("advisory", []), key=lambda r: r["_start"])
    for a, b in zip(advisory, advisory[1:]):
        gap = (b["_start"] - a["_start"]).days
        if gap < ADVISORY_GAP_DAYS:
            rep.warn("w_adv_gap", days=gap, limit=ADVISORY_GAP_DAYS, a=a["name"], b=b["name"])
    finals = by_event.get("final_report", [])
    if not finals:
        return
    final_day = min(r["_start"] for r in finals)
    if advisory:
        gap = (final_day - advisory[-1]["_start"]).days
        if gap < ADVISORY_BEFORE_FINAL_DAYS:
            rep.warn("w_adv_final", days=gap, limit=ADVISORY_BEFORE_FINAL_DAYS)
    for h in by_event.get("hearing", []):
        gap = (final_day - h["_start"]).days
        if gap < HEARING_CRITICAL_DAYS:
            rep.warn("w_hear_crit", days=gap)
        elif gap < HEARING_BEFORE_FINAL_DAYS:
            rep.warn("w_hear_final", days=gap, limit=HEARING_BEFORE_FINAL_DAYS)


def contract_mm(person):
    """Contract M/M from admin documents, or allocation rate x months."""
    if person.get("contract_mm") is not None:
        return float(person["contract_mm"])
    if person.get("rate") is not None and person.get("months") is not None:
        return float(person["rate"]) * float(person["months"])
    return None


def check_mm(plan, rep, p_start, p_end):
    if not plan.get("options", {}).get("mm", True):
        return
    valid_months = set(month_keys(p_start, p_end)) if (p_start and p_end) else set()
    for person in plan.get("people", []):
        name = person.get("name", person.get("id", "?"))
        alloc = person.get("mm", {})
        got = round(sum(float(v) for v in alloc.values()), 2)
        want = contract_mm(person)
        for month, value in alloc.items():
            if float(value) > MONTH_MM_LIMIT:
                rep.warn("w_mm_month", limit=MONTH_MM_LIMIT, name=name, month=month, value=value)
            if valid_months and month not in valid_months:
                rep.warn("w_mm_outside", name=name, month=month)
        if want is None:
            continue
        rep.note("sum_mm", name=name, got=got, want=round(want, 2))
        if abs(got - want) > MM_TOLERANCE:
            rep.warn("w_mm_total", name=name, got=got, want=round(want, 2))


def check_overview(plan, rep):
    success = plan.get("overview", {}).get("success", [])
    if len(success) > MAX_SUCCESS_LINES:
        rep.warn("w_success", n=len(success), limit=MAX_SUCCESS_LINES)


def validate(plan, lang):
    rep = Report(lang)
    p_start, p_end = check_project(plan, rep)
    rows = check_tasks(plan, rep, p_start, p_end)
    buffer_info = check_buffer(plan, rows, rep, p_start, p_end)
    check_progress(plan, rows, rep, buffer_info)
    check_events(rows, rep)
    check_mm(plan, rep, p_start, p_end)
    check_overview(plan, rep)
    return rep, rows, buffer_info


def print_report(rep):
    msg = rep.msg
    print(msg["sum_head"])
    for line in rep.summary:
        print("  " + line)
    for line in rep.errors:
        print("ERROR " + line)
    for line in rep.warnings:
        print("WARN  " + line)
    if not rep.errors and not rep.warnings:
        print(msg["sum_ok"])
    print(msg["sum_counts"].format(e=len(rep.errors), w=len(rep.warnings)))


# ---------------------------------------------------------------------------
# Table rows shared by the exporters
# ---------------------------------------------------------------------------
def people_names(plan):
    return {p.get("id"): p.get("name", p.get("id")) for p in plan.get("people", [])}


def owner_text(task, names):
    owner = task.get("owner", "")
    return names.get(owner, owner)


def raci_people(plan):
    """People shown as RACI columns, in the order of the people list."""
    if not plan.get("options", {}).get("raci", False):
        return []
    return [p for p in plan.get("people", []) if p.get("id")]


def wbs_rows(plan, rows, lang):
    lab, names = LABELS[lang], people_names(plan)
    table = []
    for i, r in enumerate(rows, start=1):
        table.append([
            i, r.get("stage", ""), r.get("name", ""), lab["type"].get(r.get("type"), ""),
            r.get("output", ""), owner_text(r, names), r["_start"], r["_end"],
            days_between(r["_start"], r["_end"]), lab["status"][r["_status"]], r.get("note", ""),
        ])
    return table


# ---------------------------------------------------------------------------
# Markdown overview
# ---------------------------------------------------------------------------
def md_escape(text):
    return str(text).replace("|", "/").replace("\n", " ")


def md_table(header, body):
    lines = ["| " + " | ".join(header) + " |", "|" + "---|" * len(header)]
    for row in body:
        lines.append("| " + " | ".join(md_escape(c) for c in row) + " |")
    return lines


def md_bullets(label, items):
    items = items if isinstance(items, list) else [items]
    return [f"- ({label}) {item}" for item in items if item]


def md_wbs_section(plan, rows, lang, buffer_info):
    lab, out = LABELS[lang], []
    if buffer_info:
        out.append(lab["buffer_line"].format(
            last=buffer_info["last"], end=parse_date(plan["project"]["end"]),
            days=buffer_info["slack"], pct=pct(buffer_info["slack"], buffer_info["total"])))
        out.append("")
    stages = []
    for r in rows:
        if r.get("stage") not in stages:
            stages.append(r.get("stage"))
    # Wait and buffer rows stay in the detail table; the summary lists work and checkpoints.
    n = 0
    for stage in stages:
        members = [r for r in rows if r.get("stage") == stage]
        shown = [r for r in members if r.get("type") in ("work", "check")]
        if not shown:
            continue
        n += 1
        span = span_label(lang, min(r["_start"] for r in members), max(r["_end"] for r in members))
        out.append(f"## ({n}) {stage} ({span})")
        for r in shown:
            text = r.get("output") or r.get("note") or lab["type"][r.get("type")]
            out.append(f"- ({r.get('name')}) {text}")
        out.append("")
    return out


def md_checks_section(rows):
    out = []
    for r in sorted((r for r in rows if r.get("type") == "check"), key=lambda r: r["_start"]):
        when = str(r["_start"]) if r["_start"] == r["_end"] else f"{r['_start']} - {r['_end']}"
        out.append(f"- ({r.get('name')}) {when}")
    return out


def md_mm_section(plan, lang):
    lab = LABELS[lang]
    proj = plan["project"]
    months = month_keys(parse_date(proj["start"]), parse_date(proj["end"]))
    header = lab["mm_cols"][1:] + months + [lab["mm_sum"]]
    body = []
    for p in plan.get("people", []):
        alloc = p.get("mm", {})
        cells = [f"{float(alloc[m]):.2f}" if m in alloc else "" for m in months]
        want = contract_mm(p)
        rate = p.get("rate")
        body.append([p.get("name", ""), p.get("role", ""), "" if rate is None else f"{float(rate):.0%}",
                     "" if want is None else f"{want:.2f}"] + cells +
                    [f"{sum(float(v) for v in alloc.values()):.2f}"])
    return md_table(header, body)


def export_md(plan, rows, lang, buffer_info):
    lab, ov, proj = LABELS[lang], plan.get("overview", {}), plan["project"]
    start, end = parse_date(proj["start"]), parse_date(proj["end"])
    out = [f"# {lab['h_basic']}", f"- ({lab['l_name']}) {proj['name']}",
           f"- ({lab['l_period']}) {start} - {end} ({months_text(lang, start, end)})"]
    people = plan.get("people", [])
    if plan.get("options", {}).get("mm", True) and people:
        total_mm = sum(float(v) for p in people for v in p.get("mm", {}).values())
        out.append(f"- ({lab['l_staff']}) " + lab["staff_line"].format(n=len(people), mm=f"{total_mm:.1f}"))
    for label, value in ov.get("summary", []):
        out.append(f"- ({label}) {value}")
    out += ["", f"# {lab['h_why']}"]
    out += md_bullets(lab["l_benef"], ov.get("beneficiary", []))
    out += md_bullets(lab["l_needs"], ov.get("needs", []))
    out += ["", f"# {lab['h_goals']}"]
    out += [f"- ({lab['l_goal'].format(n=n)}) {g}" for n, g in enumerate(ov.get("goals", []), 1)]
    out += ["", f"# {lab['h_wbs']}"] + md_wbs_section(plan, rows, lang, buffer_info)
    out += [f"# {lab['h_hyp']}"]
    for n, h in enumerate(ov.get("hypotheses", []), 1):
        out.append(f"## {lab['l_hyp'].format(n=n, text=h.get('statement', ''))}")
        out += [f"- {s}" for s in h.get("solutions", [])]
        out.append("")
    out += [f"# {lab['h_success']}"] + [f"- {s}" for s in ov.get("success", [])]
    out += ["", f"# {lab['h_checks']}"] + md_checks_section(rows)
    out += ["", f"# {lab['h_notes']}"] + [f"- {s}" for s in ov.get("notes", [])]
    header = list(lab["wbs_cols"])
    body = wbs_rows(plan, rows, lang)
    raci_cols = raci_people(plan)
    if raci_cols:
        header.append("RACI")
        for row, r in zip(body, rows):
            pairs = [f"{p['name']} {r.get('raci', {}).get(p['id'])}"
                     for p in raci_cols if r.get("raci", {}).get(p["id"])]
            row.append(", ".join(pairs))
    out += ["", f"# {lab['h_detail']}"] + md_table(header, body)
    if plan.get("options", {}).get("mm", True) and plan.get("people"):
        out += ["", f"# {lab['h_mm']}"] + md_mm_section(plan, lang)
    return "\n".join(out) + "\n"


# ---------------------------------------------------------------------------
# CSV (UTF-8 with BOM so that spreadsheet programs read Korean correctly)
# ---------------------------------------------------------------------------
def export_csv(plan, rows, lang, path):
    lab = LABELS[lang]
    raci_cols = raci_people(plan)
    header = lab["wbs_cols"] + [f"{p['name']} (RACI)" for p in raci_cols]
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(header)
        for row, r in zip(wbs_rows(plan, rows, lang), rows):
            row = [c.isoformat() if isinstance(c, date) else c for c in row]
            writer.writerow(row + [r.get("raci", {}).get(p["id"], "") for p in raci_cols])
    written = [path]
    if plan.get("options", {}).get("mm", True) and plan.get("people"):
        mm_path = path[:-4] + "_mm.csv" if path.lower().endswith(".csv") else path + "_mm.csv"
        proj = plan["project"]
        months = month_keys(parse_date(proj["start"]), parse_date(proj["end"]))
        with open(mm_path, "w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(lab["mm_cols"] + months + [lab["mm_sum"]])
            for n, p in enumerate(plan["people"], start=1):
                alloc = p.get("mm", {})
                want = contract_mm(p)
                writer.writerow([n, p.get("name", ""), p.get("role", ""), p.get("rate", ""),
                                 "" if want is None else round(want, 2)] +
                                [alloc.get(m, "") for m in months] +
                                [round(sum(float(v) for v in alloc.values()), 2)])
        written.append(mm_path)
    return written


# ---------------------------------------------------------------------------
# XLSX: plain formatting, black text only, formulas keep the Gantt live
# ---------------------------------------------------------------------------
HEADER_FILL = "D9D9D9"   # a standard Excel theme gray that is easy to pick again
TITLE_ROWS = 3           # title, period, legend
LABEL_ROW = TITLE_ROWS + 1   # month (weekly) or year (monthly) labels above the Gantt dates
COL_WIDTH = {"No": 5, "stage": 14, "task": 30, "type": 7, "output": 34, "owner": 12,
             "date": 11, "days": 7, "status": 7, "note": 18, "raci": 6,
             "gantt_week": 4, "gantt_month": 5.5}


def gantt_periods(start, end):
    """Weekly columns for short projects, monthly columns otherwise."""
    weeks = (end - start).days // 7 + 1
    if weeks <= WEEKLY_GANTT_MAX_WEEKS:
        monday = start - timedelta(days=start.weekday())
        periods = []
        while monday <= end:
            periods.append(monday)
            monday += timedelta(days=7)
        return "week", periods
    return "month", [date(int(k[:4]), int(k[5:]), 1) for k in month_keys(start, end)]


def gantt_formula(unit, row, head_ref, start_ref, end_ref, type_ref, lab):
    t = lab["type"]
    symbol = (f'IF({type_ref}="{t["wait"]}","{SYMBOL["wait"]}",'
              f'IF({type_ref}="{t["check"]}","{SYMBOL["check"]}",'
              f'IF({type_ref}="{t["buffer"]}","{SYMBOL["buffer"]}","{SYMBOL["work"]}")))')
    period_end = f"{head_ref}+6" if unit == "week" else f"EOMONTH({head_ref},0)"
    return (f'=IF(OR({start_ref}="",{end_ref}=""),"",'
            f'IF(AND({start_ref}<={period_end},{end_ref}>={head_ref}),{symbol},""))')


def style_header(ws, row, first_col, last_col, styles):
    for col in range(first_col, last_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.font, cell.fill, cell.border = styles["bold"], styles["fill"], styles["border"]
        cell.alignment = styles["center_wrap"]


def make_styles(lang):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    thin = Side(style="thin", color="000000")
    name = LABELS[lang]["font"]
    return {
        "font": Font(name=name, size=10, color="000000"),
        "bold": Font(name=name, size=10, bold=True, color="000000"),
        "title": Font(name=name, size=12, bold=True, color="000000"),
        "fill": PatternFill("solid", fgColor=HEADER_FILL),
        "border": Border(left=thin, right=thin, top=thin, bottom=thin),
        "wrap": Alignment(vertical="center", wrap_text=True),
        "center": Alignment(horizontal="center", vertical="center"),
        "center_wrap": Alignment(horizontal="center", vertical="center", wrap_text=True),
        "label": Alignment(horizontal="center", vertical="center", shrink_to_fit=True),
        "white": PatternFill("solid", fgColor="FFFFFF"),
    }


def write_gantt_labels(ws, unit, periods, first_col, lang, styles):
    """Label row above the Gantt dates: one merged cell per month over weekly columns,
    or one merged cell per year over monthly columns, so each group reads as a block."""
    from openpyxl.utils import get_column_letter
    lab = LABELS[lang]
    groups = []  # [key, first index, last index]
    for k, day in enumerate(periods):
        key = (day.year, day.month) if unit == "week" else day.year
        if groups and groups[-1][0] == key:
            groups[-1][2] = k
        else:
            groups.append([key, k, k])
    for key, first, last in groups:
        day = periods[first]
        template = lab["week_label"] if unit == "week" else lab["year_label"]
        text = template.format(y=day.year, m=day.month, mon=MONTH_ABBR[day.month - 1])
        for k in range(first, last + 1):
            cell = ws.cell(row=LABEL_ROW, column=first_col + k)
            cell.font, cell.fill, cell.border = styles["bold"], styles["white"], styles["border"]
            cell.alignment = styles["label"]
        ws.cell(row=LABEL_ROW, column=first_col + first, value=text)
        if last > first:
            ws.merge_cells(f"{get_column_letter(first_col + first)}{LABEL_ROW}:"
                           f"{get_column_letter(first_col + last)}{LABEL_ROW}")


def write_wbs_sheet(wb, plan, rows, lang, styles):
    from openpyxl.utils import get_column_letter
    lab, proj = LABELS[lang], plan["project"]
    p_start, p_end = parse_date(proj["start"]), parse_date(proj["end"])
    ws = wb.active
    ws.title = lab["sheet_wbs"]
    ws["A1"], ws["A1"].font = proj["name"], styles["title"]
    ws["A2"], ws["A2"].font = f"{lab['period']} {p_start} - {p_end}", styles["font"]
    raci_cols = raci_people(plan)
    header = lab["wbs_cols"] + [f"{p['name']}\n(RACI)" for p in raci_cols]
    unit, periods = gantt_periods(p_start, p_end)
    legend = lab["legend"] + (lab["week_note"] if unit == "week" else "")
    ws["A3"], ws["A3"].font = legend, styles["font"]
    head_row = LABEL_ROW + 1
    gantt_first = len(header) + 1
    for col, text in enumerate(header, start=1):
        ws.cell(row=head_row, column=col, value=text)
    for k, day in enumerate(periods):
        cell = ws.cell(row=head_row, column=gantt_first + k, value=day)
        cell.number_format = "d" if unit == "week" else lab["month_fmt"]
    style_header(ws, head_row, 1, gantt_first + len(periods) - 1, styles)
    write_gantt_labels(ws, unit, periods, gantt_first, lang, styles)

    idx = {name: i + 1 for i, name in enumerate(
        ["no", "stage", "task", "type", "output", "owner", "start", "end", "days", "status", "note"])}
    letter = {k: get_column_letter(v) for k, v in idx.items()}
    for n, (values, r) in enumerate(zip(wbs_rows(plan, rows, lang), rows)):
        row = head_row + 1 + n
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.font, cell.border, cell.alignment = styles["font"], styles["border"], styles["wrap"]
        for key in ("start", "end"):
            ws.cell(row=row, column=idx[key]).number_format = "yyyy-mm-dd"
        ws.cell(row=row, column=idx["days"]).value = (
            f'=IF(OR({letter["start"]}{row}="",{letter["end"]}{row}=""),"",'
            f'{letter["end"]}{row}-{letter["start"]}{row}+1)')
        for key in ("no", "type", "start", "end", "days", "status"):
            ws.cell(row=row, column=idx[key]).alignment = styles["center"]
        for j, person in enumerate(raci_cols):
            cell = ws.cell(row=row, column=len(lab["wbs_cols"]) + 1 + j,
                           value=r.get("raci", {}).get(person["id"], ""))
            cell.font, cell.border, cell.alignment = styles["font"], styles["border"], styles["center"]
        for k in range(len(periods)):
            col = gantt_first + k
            head_ref = f"{get_column_letter(col)}${head_row}"
            cell = ws.cell(row=row, column=col, value=gantt_formula(
                unit, row, head_ref, f"${letter['start']}{row}", f"${letter['end']}{row}",
                f"${letter['type']}{row}", lab))
            cell.font, cell.border, cell.alignment = styles["font"], styles["border"], styles["center"]

    widths = [COL_WIDTH["No"], COL_WIDTH["stage"], COL_WIDTH["task"], COL_WIDTH["type"],
              COL_WIDTH["output"], COL_WIDTH["owner"], COL_WIDTH["date"], COL_WIDTH["date"],
              COL_WIDTH["days"], COL_WIDTH["status"], COL_WIDTH["note"]]
    gantt_width = COL_WIDTH["gantt_week"] if unit == "week" else COL_WIDTH["gantt_month"]
    widths += [COL_WIDTH["raci"]] * len(raci_cols) + [gantt_width] * len(periods)
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    last_row = head_row + len(rows)
    ws.freeze_panes = ws.cell(row=head_row + 1, column=idx["output"])
    ws.auto_filter.ref = f"A{head_row}:{get_column_letter(len(header))}{last_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def write_mm_sheet(wb, plan, lang, styles):
    from openpyxl.utils import get_column_letter
    lab, proj = LABELS[lang], plan["project"]
    months = month_keys(parse_date(proj["start"]), parse_date(proj["end"]))
    ws = wb.create_sheet(lab["sheet_mm"])
    ws["A1"], ws["A1"].font = proj["name"], styles["title"]
    head_row = 3
    header = lab["mm_cols"] + months + [lab["mm_sum"], lab["mm_diff"]]
    for col, text in enumerate(header, start=1):
        ws.cell(row=head_row, column=col, value=text)
    style_header(ws, head_row, 1, len(header), styles)
    first_m = len(lab["mm_cols"]) + 1
    last_m = first_m + len(months) - 1
    sum_col, diff_col = last_m + 1, last_m + 2
    people = plan.get("people", [])
    for n, p in enumerate(people, start=1):
        row = head_row + n
        want = contract_mm(p)
        values = [n, p.get("name", ""), p.get("role", ""), p.get("rate"),
                  None if want is None else round(want, 2)]
        values += [p.get("mm", {}).get(m) for m in months]
        for col, value in enumerate(values, start=1):
            ws.cell(row=row, column=col, value=value)
        fl, ll = get_column_letter(first_m), get_column_letter(last_m)
        ws.cell(row=row, column=sum_col, value=f"=SUM({fl}{row}:{ll}{row})")
        ws.cell(row=row, column=diff_col, value=f'=IF(E{row}="","",{get_column_letter(sum_col)}{row}-E{row})')
    total_row = head_row + len(people) + 1
    ws.cell(row=total_row, column=2, value=lab["mm_month_sum"])
    for col in list(range(first_m, last_m + 1)) + [sum_col]:
        cl = get_column_letter(col)
        ws.cell(row=total_row, column=col, value=f"=SUM({cl}{head_row + 1}:{cl}{total_row - 1})")
    for row in range(head_row + 1, total_row + 1):
        for col in range(1, len(header) + 1):
            cell = ws.cell(row=row, column=col)
            cell.font, cell.border = styles["font"], styles["border"]
            if col >= 4:
                cell.number_format = "0%" if col == 4 else "0.00"
                cell.alignment = styles["center"]
        ws.cell(row=row, column=1).alignment = styles["center"]
    ws.cell(row=total_row, column=2).font = styles["bold"]
    widths = [5, 14, 14, 8, 9] + [8] * len(months) + [8, 8]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = ws.cell(row=head_row + 1, column=first_m)
    ws.page_setup.orientation = "landscape"


def export_xlsx(plan, rows, lang, path):
    try:
        from openpyxl import Workbook
    except ImportError:
        sys.exit("openpyxl is required: pip install openpyxl")
    wb = Workbook()
    styles = make_styles(lang)
    write_wbs_sheet(wb, plan, rows, lang, styles)
    if plan.get("options", {}).get("mm", True) and plan.get("people"):
        write_mm_sheet(wb, plan, lang, styles)
    wb.save(path)
    return [path]


# ---------------------------------------------------------------------------
# Command line
# ---------------------------------------------------------------------------
def main():
    parser = argparse.ArgumentParser(description="Check and export a project plan.")
    parser.add_argument("command", choices=["check", "md", "csv", "xlsx"])
    parser.add_argument("plan", help="path to plan.json")
    parser.add_argument("-o", "--output", help="output file path")
    parser.add_argument("--lang", default=DEFAULT_LANG, choices=sorted(LABELS))
    args = parser.parse_args()

    plan = load_plan(args.plan)
    rep, rows, buffer_info = validate(plan, args.lang)
    print_report(rep)
    if rep.errors:
        sys.exit(1)
    if args.command == "check":
        return
    default_out = {"md": "overview.md", "csv": "wbs.csv", "xlsx": "wbs.xlsx"}[args.command]
    out = args.output or default_out
    if args.command == "md":
        with open(out, "w", encoding="utf-8") as f:
            f.write(export_md(plan, rows, args.lang, buffer_info))
        written = [out]
    elif args.command == "csv":
        written = export_csv(plan, rows, args.lang, out)
    else:
        written = export_xlsx(plan, rows, args.lang, out)
    for path in written:
        print("OK " + path)


if __name__ == "__main__":
    main()
