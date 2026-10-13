"""Shared helpers for practical-calc-lite scripts: input loading, validation, and output.

Every calculator reads one JSON file (or stdin with "-") and prints a
Markdown report. Every table in the report can also be saved as CSV files
(--csv DIR) or as an Excel workbook (--xlsx PATH). Keeping I/O here lets each
calculator hold only its formulas.
"""

import argparse
import csv
import json
import math
import re
import sys
import unicodedata
from pathlib import Path

DECIMALS_DEFAULT = 2
CSV_ENCODING = "utf-8-sig"        # BOM so spreadsheet programs read Korean text correctly
SHEET_NAME_MAX = 31               # Excel limit
COLUMN_WIDTH_MAX = 40
SHEET_FORBIDDEN = r"[\[\]:*?/\\]"

_EXPORT = {"args": None, "input": None, "tables": [], "sheet_writers": []}


class InputError(ValueError):
    """Raised when user input is missing or invalid. Printed without traceback."""


def build_parser(description, example):
    """Return an argument parser with the options shared by all calculators."""
    parser = argparse.ArgumentParser(description=description)
    parser.add_argument("input", nargs="?", help="input JSON path, or - for stdin")
    parser.add_argument("--example", action="store_true",
                        help="print an example input JSON and exit")
    parser.add_argument("--json", metavar="PATH",
                        help="also write machine-readable results to PATH")
    parser.add_argument("--xlsx", metavar="PATH",
                        help="also save every report table as a sheet of an Excel workbook")
    parser.add_argument("--csv", metavar="DIR",
                        help="also save every report table as a CSV file in DIR")
    parser.set_defaults(_example=example)
    return parser


def load_input(args):
    """Load the JSON input named on the command line."""
    if args.example:
        print(json.dumps(args._example, ensure_ascii=False, indent=2))
        sys.exit(0)
    if not args.input:
        raise InputError("input JSON path is required (use --example for a template)")
    if args.input == "-":
        data = json.load(sys.stdin)
    else:
        path = Path(args.input)
        if not path.is_file():
            raise InputError(f"input file not found: {path}")
        with path.open(encoding="utf-8") as handle:
            data = json.load(handle)
    _EXPORT["args"], _EXPORT["input"] = args, data
    return data


def require(data, key, where="input"):
    """Return data[key] or raise a readable error naming the missing field."""
    if key not in data or data[key] is None:
        raise InputError(f"missing required field '{key}' in {where}")
    return data[key]


NEGATIVE_MARK = "△"               # Korean tables write negatives as "△100"
AMBIGUOUS_MARKS = ("▲", "▼", "▽")  # meaning differs by publisher; convert during preparation
MINUS_VARIANTS = ("−", "–")       # unicode minus and en dash typed as minus
MISSING_MARKS = {"", "-", "…", "...", "x", "n/a", "na", "nan"}  # compared in lower case


def parse_number_text(text, name):
    """Read common spreadsheet notations: 1,234 / (100) / △100 / 4.5% / full-width digits.

    Unit words such as "억원" are rejected on purpose: unit conversion belongs to the
    data-preparation step, where the factor is written down.
    """
    raw = unicodedata.normalize("NFKC", text).strip().replace(",", "").replace(" ", "")
    for variant in MINUS_VARIANTS:
        raw = raw.replace(variant, "-")
    if raw.lower() in MISSING_MARKS:
        raise InputError(f"'{name}': {text!r} marks a blank or suppressed cell; decide during "
                         "preparation whether it means zero or missing")
    if any(mark in raw for mark in AMBIGUOUS_MARKS):
        raise InputError(f"'{name}': {text!r} uses a mark whose sign depends on the source; "
                         "convert it to a plain number first")
    sign = 1.0
    if raw.startswith(NEGATIVE_MARK):
        sign, raw = -1.0, raw[len(NEGATIVE_MARK):]
    if raw.startswith("(") and raw.endswith(")"):
        sign, raw = -sign, raw[1:-1]
    scale = 1.0
    if raw.endswith("%"):
        scale, raw = 0.01, raw[:-1]
    try:
        return sign * float(raw) * scale
    except ValueError as exc:
        raise InputError(f"'{name}' must be a number, got {text!r} "
                         "(remove unit words and convert units during preparation)") from exc


def to_number(value, name):
    """Convert value to float and reject NaN or infinity."""
    if isinstance(value, str):
        number = parse_number_text(value, name)
    else:
        try:
            number = float(value)
        except (TypeError, ValueError) as exc:
            raise InputError(f"'{name}' must be a number, got {value!r}") from exc
    if math.isnan(number) or math.isinf(number):
        raise InputError(f"'{name}' must be finite, got {value!r}")
    return number


def fmt(value, decimals=DECIMALS_DEFAULT):
    """Format numbers with thousands separators; pass other values through."""
    if value is None:
        return "-"
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return f"{value:,}"
    if isinstance(value, float):
        if math.isnan(value) or math.isinf(value):
            return "n/a"
        return f"{value:,.{decimals}f}"
    return str(value)


def pct(value, decimals=2):
    """Format a ratio (0.1234) as a percentage string (12.34%)."""
    if value is None:
        return "-"
    return f"{value * 100:,.{decimals}f}%"


class Table(str):
    """Markdown table text that also keeps the raw headers and rows for file export."""

    def __new__(cls, text, headers, rows, decimals):
        obj = super().__new__(cls, text)
        obj.headers, obj.rows, obj.decimals = list(headers), [list(r) for r in rows], decimals
        return obj


def md_table(headers, rows, decimals=DECIMALS_DEFAULT):
    """Render a Markdown table from a header list and row lists."""
    lines = ["| " + " | ".join(headers) + " |",
             "| " + " | ".join("---" for _ in headers) + " |"]
    for row in rows:
        lines.append("| " + " | ".join(fmt(cell, decimals) for cell in row) + " |")
    return Table("\n".join(lines), headers, rows, decimals)


def section(title, body):
    """Print one report section and keep its table for --csv / --xlsx."""
    print(f"## {title}\n")
    print(body)
    print()
    if isinstance(body, Table):
        _EXPORT["tables"].append((title, body))


def register_sheet(writer):
    """Add a calculator-specific sheet; writer(workbook) runs only when --xlsx is given."""
    _EXPORT["sheet_writers"].append(writer)


def excel_cell(value, decimals):
    """Return (value, number_format) with numeric text turned back into numbers."""
    if hasattr(value, "item") and not isinstance(value, (str, bytes)):
        value = value.item()                        # numpy scalars from pandas tables
    if value is None:
        return None, None
    if isinstance(value, bool):
        return value, None
    if isinstance(value, int):
        return value, "#,##0"
    if isinstance(value, float):
        return (value, "#,##0." + "0" * decimals) if math.isfinite(value) else (None, None)
    if not isinstance(value, str):
        value = str(value)
    text = value.strip()
    if re.fullmatch(r"-?\d+", text):
        return int(text), "0"                      # years and counts printed as text
    if re.fullmatch(r"[-△(]?[\d,]+(\.\d+)?\)?%?", text):
        try:
            number = parse_number_text(text, "cell")
        except InputError:
            return value, None
        places = len(text.rstrip("%)").split(".")[1]) if "." in text else 0
        if text.endswith("%"):
            return number, "0." + "0" * places + "%" if places else "0%"
        return number, "#,##0." + "0" * places if places else "#,##0"
    return value, None


def unique_sheet_title(index, title, used):
    base = re.sub(SHEET_FORBIDDEN, " ", f"{index:02d} {title}")[:SHEET_NAME_MAX].strip()
    name, n = base, 2
    while name in used:
        suffix = f" ({n})"
        name, n = base[:SHEET_NAME_MAX - len(suffix)] + suffix, n + 1
    used.add(name)
    return name


def write_csv_tables(directory):
    out = Path(directory)
    out.mkdir(parents=True, exist_ok=True)
    for index, (title, table) in enumerate(_EXPORT["tables"], 1):
        slug = re.sub(r"[^0-9A-Za-z가-힣]+", "_", title).strip("_")[:60] or "table"
        path = out / f"{index:02d}_{slug}.csv"
        with path.open("w", encoding=CSV_ENCODING, newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(table.headers)
            for row in table.rows:
                writer.writerow(["" if cell is None else cell for cell in row])
    print(f"(tables written to {out}/ as {len(_EXPORT['tables'])} CSV files)")


def write_xlsx_tables(path):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    workbook.remove(workbook.active)
    used = set()
    for index, (title, table) in enumerate(_EXPORT["tables"], 1):
        sheet = workbook.create_sheet(unique_sheet_title(index, title, used))
        sheet.append([title])
        sheet["A1"].font = Font(bold=True)
        sheet.append(table.headers)
        for cell in sheet[2]:
            cell.font = Font(bold=True)
        for row in table.rows:
            converted = [excel_cell(v, table.decimals) for v in row]
            sheet.append([v for v, _ in converted])
            for cell, (_, number_format) in zip(sheet[sheet.max_row], converted):
                if number_format:
                    cell.number_format = number_format
        sheet.freeze_panes = "A3"
        for col in range(1, len(table.headers) + 1):
            width = max(len(str(sheet.cell(r, col).value or "")) for r in range(2, sheet.max_row + 1))
            sheet.column_dimensions[get_column_letter(col)].width = min(COLUMN_WIDTH_MAX, width + 2)
    for writer in _EXPORT["sheet_writers"]:
        writer(workbook)
    inputs = workbook.create_sheet("input")
    inputs.append(["Input JSON used for this workbook (rerun the script with it to reproduce)"])
    for line in json.dumps(_EXPORT["input"], ensure_ascii=False, indent=2).splitlines():
        inputs.append([line])
    workbook.save(path)
    print(f"(tables written to {path})")


def export_tables():
    """Write --csv / --xlsx outputs after a successful run."""
    args = _EXPORT["args"]
    if args is None:
        return
    if getattr(args, "csv", None):
        write_csv_tables(args.csv)
    if getattr(args, "xlsx", None):
        write_xlsx_tables(args.xlsx)


def write_json(path, payload):
    """Write results for downstream use (charts, spreadsheets, re-checks)."""
    if not path:
        return
    with open(path, "w", encoding="utf-8") as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=2)
    print(f"(results written to {path})")


def run(main_func):
    """Run a calculator and turn InputError into a clean message and exit code 2."""
    try:
        main_func()
        export_tables()
    except InputError as exc:
        print(f"[INPUT ERROR] {exc}", file=sys.stderr)
        sys.exit(2)
